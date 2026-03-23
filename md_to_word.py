import re
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def markdown_to_docx(md_file, docx_file):
    document = Document()
    
    # 设置默认字体为宋体/Times New Roman
    style = document.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    document.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), u'宋体')

    with open(md_file, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    code_mode = False
    
    for line in lines:
        line = line.rstrip()
        
        # 处理代码块
        if line.startswith('```'):
            code_mode = not code_mode
            continue
        
        if code_mode:
            p = document.add_paragraph()
            runner = p.add_run(line)
            runner.font.name = 'Courier New'
            runner.font.size = Pt(10)
            runner.font.color.rgb = RGBColor(100, 100, 100)
            p.paragraph_format.left_indent = Pt(20)
            continue

        # 处理标题
        if line.startswith('# '):
            document.add_heading(line[2:], level=0)
        elif line.startswith('## '):
            document.add_heading(line[3:], level=1)
        elif line.startswith('### '):
            document.add_heading(line[4:], level=2)
        elif line.startswith('#### '):
            document.add_heading(line[5:], level=3)
        
        # 处理分隔线
        elif line.startswith('---'):
            document.add_page_break()
            
        # 处理列表
        elif line.strip().startswith('* ') or line.strip().startswith('- '):
            p = document.add_paragraph(line.strip()[2:], style='List Bullet')
            
        elif re.match(r'^\d+\.\s', line.strip()):
            content = re.sub(r'^\d+\.\s', '', line.strip())
            p = document.add_paragraph(content, style='List Number')
            
        # 处理普通文本
        else:
            if not line:
                continue
            
            p = document.add_paragraph()
            
            # 简单的加粗解析 (**text**)
            parts = re.split(r'(\*\*.*?\*\*)', line)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    p.add_run(part)

    document.save(docx_file)
    print(f"Successfully created {docx_file}")

if __name__ == "__main__":
    markdown_to_docx("project_report.md", "项目报告_完整版.docx")

